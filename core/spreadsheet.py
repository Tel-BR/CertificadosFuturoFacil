from datetime import datetime
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

from core.validator import ValidacaoAluno, validar_aluno


MESES_PT = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}

# Regex oficial: (Nh) [I|II|III] AAAA/mmm/DD
ENCOUNTER_REGEX = re.compile(
    r"^\s*\((?P<horas>\d+)h\)\s*(?:(?P<periodo>[IVXLCDM]+)\s+)?(?P<ano>\d{4})/(?P<mes>[a-zA-Z]{3})/(?P<dia>\d{1,2})\s*$",
    re.IGNORECASE,
)


@dataclass
class EncontroColuna:
    """Metadados de uma coluna de encontro/frequência identificada na planilha."""

    nome_coluna: str
    horas: int
    periodo: Optional[str]
    data: datetime
    data_str: str


@dataclass
class SpreadsheetValidationResult:
    """Resultado da auditoria e validação de uma planilha de alunos."""

    total_rows: int
    valid_count: int
    invalid_count: int
    is_valid: bool
    sem_cpf_count: int = 0
    alunos: List[ValidacaoAluno] = field(default_factory=list)
    global_errors: List[str] = field(default_factory=list)
    carga_horaria_calculada: Optional[int] = None
    data_inicio_calculada: Optional[str] = None
    data_fim_calculada: Optional[str] = None
    encontros_detectados: List[EncontroColuna] = field(default_factory=list)

    @property
    def students(self) -> List[ValidacaoAluno]:
        """Alias para compatibilidade de nomenclatura."""
        return self.alunos


MESES_NUM_TO_STR = {v: k for k, v in MESES_PT.items()}


def _parse_date_input(val: Any) -> Optional[datetime]:
    """Interpreta datas a partir de string, date ou datetime."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    from datetime import date
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def format_encounter_header(dt: Union[datetime, Any], horas: int = 4) -> str:
    """Gera o nome de coluna de encontro no padrão oficial: '(4h) AAAA/mmm/DD' sem I ou II."""
    d = _parse_date_input(dt)
    if d is None:
        raise ValueError(f"Data inválida para coluna de encontro: {dt}")
    mes_str = MESES_NUM_TO_STR.get(d.month, "jan")
    return f"({horas}h) {d.year}/{mes_str}/{d.day:02d}"


def generate_template_spreadsheet(
    destination: Union[str, Path, io.BytesIO] = "modelo_alunos.xlsx",
    data_inicio: Optional[Union[str, Any]] = None,
    data_fim: Optional[Union[str, Any]] = None,
    horas_por_encontro: int = 4,
    incluir_exemplo: bool = False,
) -> Union[Path, io.BytesIO]:
    """
    Gera a planilha modelo oficial 'modelo_alunos.xlsx' contendo as colunas 'Nome' e 'CPF'
    e opcionalmente as colunas de encontros configuradas a partir das datas da turma.
    
    Se data_inicio e/ou data_fim forem fornecidas:
    - Adiciona colunas no formato oficial '(4h) AAAA/mmm/DD' (sem numeração I/II).
    - Se incluir_exemplo=True, insere linha de exemplo com data inicial = True e data final = False.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alunos"

    # Estilos visuais para cabeçalho
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Azul Marinho
    header_align = Alignment(horizontal="center", vertical="center")

    headers = ["Nome", "CPF", "Aproveitamento (%)"]
    col_ini_hdr = None
    col_fim_hdr = None

    dt_ini = _parse_date_input(data_inicio)
    dt_fim = _parse_date_input(data_fim)

    if dt_ini:
        col_ini_hdr = format_encounter_header(dt_ini, horas=horas_por_encontro)
        headers.append(col_ini_hdr)
        if dt_fim and dt_fim.date() != dt_ini.date():
            col_fim_hdr = format_encounter_header(dt_fim, horas=horas_por_encontro)
            headers.append(col_fim_hdr)

    ws.append(headers)
    ws.row_dimensions[1].height = 25

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Ajuste de largura das colunas
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    for col_idx in range(4, len(headers) + 1):
        from openpyxl.utils import get_column_letter
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 26

    # Linha de exemplo
    if incluir_exemplo:
        exemplo_row: List[Any] = ["Maria de Souza Silva", "529.982.247-25", None]
        if col_ini_hdr:
            exemplo_row.append(True)
            if col_fim_hdr:
                exemplo_row.append(False)
        ws.append(exemplo_row)
        ws.row_dimensions[2].height = 22
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=2, column=2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=2, column=3).alignment = Alignment(horizontal="center", vertical="center")
        if col_ini_hdr:
            ws.cell(row=2, column=4).alignment = Alignment(horizontal="center", vertical="center")
            if col_fim_hdr:
                ws.cell(row=2, column=5).alignment = Alignment(horizontal="center", vertical="center")

    # Salvar para caminho ou buffer
    if isinstance(destination, io.IOBase):
        wb.save(destination)
        return destination
    else:
        path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return path


def _carregar_dataframe(source: Union[str, Path, io.BytesIO, bytes, pd.DataFrame]) -> pd.DataFrame:
    """Carrega dados tabulares de arquivo Excel ou CSV (com suporte a vírgula ou ponto-e-vírgula)."""
    if isinstance(source, pd.DataFrame):
        return source.copy()

    if isinstance(source, bytes):
        buffer = io.BytesIO(source)
        try:
            return pd.read_excel(buffer)
        except Exception:
            buffer.seek(0)
            return pd.read_csv(buffer, sep=None, engine="python")

    if isinstance(source, (str, Path)):
        path_str = str(source).lower()
        if path_str.endswith(".csv"):
            return pd.read_csv(source, sep=None, engine="python")
        return pd.read_excel(source)

    if isinstance(source, io.IOBase):
        try:
            return pd.read_excel(source)
        except Exception:
            if hasattr(source, "seek"):
                source.seek(0)
            return pd.read_csv(source, sep=None, engine="python")

    raise ValueError("Fonte de dados inválida ou formato não reconhecido.")


def read_and_validate_spreadsheet(
    source: Union[str, Path, io.BytesIO, bytes, pd.DataFrame],
    cpf_obrigatorio: bool = False,
    frequencia_minima: int = 75,
) -> SpreadsheetValidationResult:
    """
    Lê e audita uma planilha (Excel ou CSV), verificando a presença obrigatória
    das colunas 'Nome' e 'CPF' e suportando colunas dinâmicas de frequência/encontros
    no formato '(Nh) [I|II|III] AAAA/mmm/DD'.
    Quando cpf_obrigatorio=False, aceita alunos com CPF em branco.
    Aplica a regra de frequência mínima (padrão 75%).
    """
    global_errors: List[str] = []

    try:
        df = _carregar_dataframe(source)
    except Exception as e:
        return SpreadsheetValidationResult(
            total_rows=0,
            valid_count=0,
            invalid_count=0,
            is_valid=False,
            global_errors=[f"Erro ao ler arquivo: {str(e)}"],
        )

    col_map = {}
    encontros: List[EncontroColuna] = []
    colunas_extras = []

    for col in df.columns:
        col_str = str(col).strip()
        clean_col = col_str.lower()
        if clean_col == "nome":
            col_map[col] = "Nome"
        elif clean_col == "cpf":
            col_map[col] = "CPF"
        elif clean_col in (
            "frequencia", "frequência", "freq", "frequencia (%)", "frequência (%)",
            "aproveitamento", "aproveitamento (%)", "aprov", "aprov (%)",
        ):
            col_map[col] = "Frequencia"
        else:
            m = ENCOUNTER_REGEX.match(col_str)
            if m:
                horas = int(m.group("horas"))
                periodo = m.group("periodo")
                ano = int(m.group("ano"))
                mes_str = m.group("mes").lower()
                mes = MESES_PT.get(mes_str)
                dia = int(m.group("dia"))
                if mes and 1 <= dia <= 31:
                    try:
                        dt = datetime(ano, mes, dia)
                        encontros.append(
                            EncontroColuna(
                                nome_coluna=col,
                                horas=horas,
                                periodo=periodo,
                                data=dt,
                                data_str=dt.strftime("%d/%m/%Y"),
                            )
                        )
                    except ValueError:
                        colunas_extras.append(col_str)
                else:
                    colunas_extras.append(col_str)
            else:
                colunas_extras.append(col_str)

    if colunas_extras:
        global_errors.append(
            f"A planilha contém colunas não permitidas: {', '.join(colunas_extras)}. "
            "As colunas permitidas são 'Nome', 'CPF', 'Aproveitamento (%)' e encontros no formato '(4h) AAAA/mmm/DD'."
        )

    if "Nome" not in col_map.values() or "CPF" not in col_map.values():
        missing = []
        if "Nome" not in col_map.values():
            missing.append("'Nome'")
        if "CPF" not in col_map.values():
            missing.append("'CPF'")
        global_errors.append(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}. "
            "A planilha deve conter estritamente as colunas 'Nome' e 'CPF'."
        )

    if global_errors:
        return SpreadsheetValidationResult(
            total_rows=len(df),
            valid_count=0,
            invalid_count=len(df),
            is_valid=False,
            global_errors=global_errors,
        )

    # Localizar nomes reais das colunas
    nome_col = [orig for orig, mapped in col_map.items() if mapped == "Nome"][0]
    cpf_col = [orig for orig, mapped in col_map.items() if mapped == "CPF"][0]

    # Ordenação cronológica dos encontros identificados
    encontros.sort(key=lambda e: (e.data, e.periodo or ""))
    carga_horaria_calc = sum(e.horas for e in encontros) if encontros else None
    data_inicio_calc = encontros[0].data_str if encontros else None
    data_fim_calc = encontros[-1].data_str if encontros else None

    alunos: List[ValidacaoAluno] = []
    for _, row in df.iterrows():
        raw_name = str(row[nome_col]) if pd.notna(row[nome_col]) else ""
        val_cpf = row[cpf_col]

        # Trata inteiros do pandas onde o zero à esquerda foi suprimido pelo Excel
        if isinstance(val_cpf, (int, float)) and pd.notna(val_cpf):
            str_cpf = str(int(val_cpf))
            raw_cpf = str_cpf.zfill(11) if len(str_cpf) == 10 else str_cpf
        elif pd.notna(val_cpf):
            raw_cpf = str(val_cpf)
        else:
            raw_cpf = ""

        # Apuração de frequência / aproveitamento
        horas_presentes = None
        freq_calc = None

        # 1. Verifica se o usuário informou diretamente o número de aproveitamento/frequência
        if "Frequencia" in col_map.values():
            freq_col_name = [orig for orig, mapped in col_map.items() if mapped == "Frequencia"][0]
            val_freq = row[freq_col_name]
            if pd.notna(val_freq):
                val_freq_clean = str(val_freq).replace("%", "").strip()
                if val_freq_clean != "":
                    try:
                        freq_calc = round(float(val_freq_clean))
                    except ValueError:
                        freq_calc = None

        # 2. Se o número NÃO foi fornecido manualmente (vazio ou em branco),
        # calcula normalmente a partir dos booleanos dos encontros
        if freq_calc is None:
            if encontros and carga_horaria_calc and carga_horaria_calc > 0:
                horas_presentes = 0
                for enc in encontros:
                    cell_val = row[enc.nome_coluna]
                    is_presente = False
                    if isinstance(cell_val, bool):
                        is_presente = cell_val
                    elif pd.notna(cell_val):
                        val_s = str(cell_val).strip().lower()
                        if val_s in ("true", "1", "1.0", "verdadeiro", "v", "sim", "s", "x"):
                            is_presente = True
                    if is_presente:
                        horas_presentes += enc.horas

                freq_calc = round((horas_presentes / carga_horaria_calc) * 100)
            else:
                freq_calc = 100
        else:
            # Número fornecido manualmente: NÃO precisa calcular, é a fonte da verdade!
            if carga_horaria_calc and carga_horaria_calc > 0:
                horas_presentes = round((freq_calc / 100) * carga_horaria_calc)

        aluno = validar_aluno(
            raw_name,
            raw_cpf,
            cpf_obrigatorio=cpf_obrigatorio,
            frequencia=freq_calc,
            horas_presentes=horas_presentes,
            horas_totais=carga_horaria_calc,
            frequencia_minima=frequencia_minima,
        )
        alunos.append(aluno)

    total_rows = len(alunos)
    valid_count = sum(1 for a in alunos if a.is_valido)
    sem_cpf_count = sum(1 for a in alunos if a.is_valido and not a.cpf)
    invalid_count = total_rows - valid_count
    is_valid = (invalid_count == 0) and (total_rows > 0)

    if total_rows == 0:
        global_errors.append("A planilha está vazia.")
        is_valid = False

    return SpreadsheetValidationResult(
        total_rows=total_rows,
        valid_count=valid_count,
        invalid_count=invalid_count,
        is_valid=is_valid,
        sem_cpf_count=sem_cpf_count,
        alunos=alunos,
        global_errors=global_errors,
        carga_horaria_calculada=carga_horaria_calc,
        data_inicio_calculada=data_inicio_calc,
        data_fim_calculada=data_fim_calc,
        encontros_detectados=encontros,
    )
