"""Testes unitários e de integração para o Livro de Registro Digital."""

import os
from pathlib import Path
import openpyxl
import pytest

from core.validator import ValidacaoAluno, validar_aluno


def test_import_registry_module():
    """Verifica que o módulo core.registry pode ser importado."""
    import core.registry as reg
    assert hasattr(reg, "LivroRegistroManager")
    assert hasattr(reg, "CertificadoRegistro")
    assert hasattr(reg, "CursoMetadata")
    assert hasattr(reg, "formatar_carga_horaria_extenso")
    assert hasattr(reg, "generate_authenticity_code")


def test_formatar_carga_horaria_extenso():
    from core.registry import formatar_carga_horaria_extenso

    assert formatar_carga_horaria_extenso(1) == "uma hora"
    assert formatar_carga_horaria_extenso(20) == "vinte horas"
    assert formatar_carga_horaria_extenso(21) == "vinte e uma horas"
    assert formatar_carga_horaria_extenso(40) == "quarenta horas"
    assert formatar_carga_horaria_extenso(51) == "cinquenta e uma horas"
    assert formatar_carga_horaria_extenso(100) == "cem horas"
    assert formatar_carga_horaria_extenso(101) == "cento e uma horas"
    assert formatar_carga_horaria_extenso(120) == "cento e vinte horas"

    with pytest.raises(ValueError):
        formatar_carga_horaria_extenso(0)

    with pytest.raises(ValueError):
        formatar_carga_horaria_extenso(-5)


def test_generate_authenticity_code():
    from core.registry import generate_authenticity_code

    code1 = generate_authenticity_code(
        aluno_cpf="12345678909",
        curso_nome="Python para Dados",
        registro_num=1,
        salt="fixed_salt_1"
    )
    code2 = generate_authenticity_code(
        aluno_cpf="12345678909",
        curso_nome="Python para Dados",
        registro_num=1,
        salt="fixed_salt_1"
    )
    code3 = generate_authenticity_code(
        aluno_cpf="98765432100",
        curso_nome="Python para Dados",
        registro_num=2,
        salt="fixed_salt_2"
    )

    # Verificação de tamanho e caracteres hexadecimais maiúsculos SHA-256
    assert len(code1) == 64
    assert code1.isupper()
    assert all(c in "0123456789ABCDEF" for c in code1)

    # Imutabilidade e determinismo com mesmo payload
    assert code1 == code2

    # Unicidade para payloads distintos
    assert code1 != code3


def test_livro_registro_sequencial_e_virada_centesimo(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "test_registros.db"
    manager = LivroRegistroManager(db_url_or_path=str(db_path), folhas_por_livro=100)

    # Inicialização: banco vazio
    next_livro, next_folha, next_reg = manager.get_next_numbers()
    assert (next_livro, next_folha, next_reg) == (1, 1, 1)

    curso = CursoMetadata(
        curso_nome="Assistente Administrativo",
        carga_horaria=40,
        data_conclusao="15/08/2026",
        data_emissao="20/08/2026",
        instrutor="Prof. Carlos Andrade",
        cidade="São Paulo - SP",
        ementa="Rotinas de escritório, redação empresarial e planilhas.",
    )

    # Inserção do primeiro registro
    aluno1 = validar_aluno("João da Silva", "12345678909")  # CPF será normalizado
    reg1 = manager.register_certificate(aluno1, curso)
    assert reg1.livro_numero == 1
    assert reg1.folha_numero == 1
    assert reg1.registro_numero == 1
    assert reg1.carga_horaria_extenso == "quarenta horas"
    assert len(reg1.codigo_autenticidade) == 64

    # Próximos números após o 1º registro
    next_livro, next_folha, next_reg = manager.get_next_numbers()
    assert (next_livro, next_folha, next_reg) == (1, 2, 2)

    # Inserção de registros até atingir a folha 100 (mais 99 alunos)
    alunos_intermediarios = [
        ValidacaoAluno(
            nome=f"Aluno {i}",
            cpf=f"11144477{i:03d}",
            cpf_formatado=f"111.444.77{i:02d}-{i%100:02d}",
            cpf_mascarado=f"***.444.77{i:02d}-**",
            is_valido=True,
        )
        for i in range(2, 101)
    ]
    batch_regs = manager.register_batch(alunos_intermediarios, curso)
    assert len(batch_regs) == 99

    # O centésimo registro deve ser Livro 1, Folha 100, Registro 100
    ultimo_lote1 = batch_regs[-1]
    assert ultimo_lote1.livro_numero == 1
    assert ultimo_lote1.folha_numero == 100
    assert ultimo_lote1.registro_numero == 100

    # Próximo deve virar para Livro 2, Folha 1, Registro 101!
    next_livro, next_folha, next_reg = manager.get_next_numbers()
    assert (next_livro, next_folha, next_reg) == (2, 1, 101)

    # Inserir o 101º registro e validar a virada
    aluno101 = ValidacaoAluno(
        nome="Aluno Centésimo Primeiro",
        cpf="99988877700",
        cpf_formatado="999.888.777-00",
        cpf_mascarado="***.888.777-**",
        is_valido=True,
    )
    reg101 = manager.register_certificate(aluno101, curso)
    assert reg101.livro_numero == 2
    assert reg101.folha_numero == 1
    assert reg101.registro_numero == 101


def test_batch_registration_spanning_multiple_books(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "multi_book.db"
    manager = LivroRegistroManager(db_url_or_path=str(db_path), folhas_por_livro=100)

    curso = CursoMetadata(
        curso_nome="Informática Básica",
        carga_horaria=20,
        data_conclusao="10/05/2026",
    )

    alunos = [
        ValidacaoAluno(
            nome=f"Estudante {i}",
            cpf=f"1234567{i:04d}",
            cpf_formatado=f"123.456.{i:03d}-00",
            cpf_mascarado=f"***.456.{i:03d}-**",
            is_valido=True,
        )
        for i in range(1, 151)
    ]

    registrados = manager.register_batch(alunos, curso)
    assert len(registrados) == 150

    # Aluno 1
    assert registrados[0].livro_numero == 1
    assert registrados[0].folha_numero == 1
    assert registrados[0].registro_numero == 1

    # Aluno 100 (fim do Livro 1)
    assert registrados[99].livro_numero == 1
    assert registrados[99].folha_numero == 100
    assert registrados[99].registro_numero == 100

    # Aluno 101 (início do Livro 2)
    assert registrados[100].livro_numero == 2
    assert registrados[100].folha_numero == 1
    assert registrados[100].registro_numero == 101

    # Aluno 150 (metade do Livro 2)
    assert registrados[149].livro_numero == 2
    assert registrados[149].folha_numero == 50
    assert registrados[149].registro_numero == 150


def test_query_certificate_by_code(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "query_test.db"
    manager = LivroRegistroManager(db_url_or_path=str(db_path))

    curso = CursoMetadata(
        curso_nome="Gestão de Pessoas",
        carga_horaria=60,
        data_conclusao="30/06/2026",
        instrutor="Dra. Helena Souza",
    )
    aluno = ValidacaoAluno(
        nome="Maria Fernandes",
        cpf="11122233344",
        cpf_formatado="111.222.333-44",
        cpf_mascarado="***.222.333-**",
        is_valido=True,
    )

    reg = manager.register_certificate(aluno, curso)
    codigo = reg.codigo_autenticidade

    # Busca exata
    encontrado = manager.get_certificate_by_code(codigo)
    assert encontrado is not None
    assert encontrado.aluno_nome == "Maria Fernandes"
    assert encontrado.aluno_cpf_mascarado == "***.222.333-**"
    assert encontrado.curso_nome == "Gestão de Pessoas"
    assert encontrado.carga_horaria == 60

    # Busca insensível a maiúsculas/minúsculas
    encontrado_lower = manager.get_certificate_by_code(codigo.lower())
    assert encontrado_lower is not None
    assert encontrado_lower.codigo_autenticidade == codigo

    # Código inexistente
    inexistente = manager.get_certificate_by_code("0" * 64)
    assert inexistente is None


def test_export_to_excel_and_auto_export(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "excel_test.db"
    excel_file = tmp_path / "livro_registro_certificados.xlsx"

    manager = LivroRegistroManager(
        db_url_or_path=str(db_path),
        excel_export_path=str(excel_file),
    )

    curso = CursoMetadata(
        curso_nome="Empreendedorismo MEI",
        carga_horaria=30,
        data_conclusao="01/09/2026",
        data_emissao="05/09/2026",
        instrutor="Roberto Silveira",
        cidade="Belo Horizonte - MG",
    )

    alunos = [
        ValidacaoAluno(
            nome=f"Empreendedor {i}",
            cpf=f"22233344{i:03d}",
            cpf_formatado=f"222.333.44{i:02d}-00",
            cpf_mascarado=f"***.333.44{i:02d}-**",
            is_valido=True,
        )
        for i in range(1, 4)
    ]

    # Registrar lote (deve disparar exportação automática para excel_file)
    manager.register_batch(alunos, curso)

    assert excel_file.exists()

    wb = openpyxl.load_workbook(str(excel_file))
    ws = wb.active
    assert ws.title == "Livro de Registro"

    # Validar cabeçalho
    expected_headers = [
        "Livro Nº",
        "Folha Nº",
        "Registro Nº",
        "Código de Autenticidade",
        "Nome do Aluno",
        "CPF",
        "Curso",
        "Carga Horária",
        "Data de Início",
        "Data de Conclusão",
        "Data de Emissão",
        "Modalidade",
        "Instrutor",
        "Cidade",
    ]
    actual_headers = [ws.cell(row=1, column=col).value for col in range(1, len(expected_headers) + 1)]
    assert actual_headers == expected_headers

    # Validar número de linhas (cabeçalho + 3 registros)
    assert ws.max_row == 4

    # Validar dados da 1ª linha de dados (linha 2)
    assert ws.cell(row=2, column=1).value == 1  # Livro
    assert ws.cell(row=2, column=2).value == 1  # Folha
    assert ws.cell(row=2, column=3).value == 1  # Registro
    assert ws.cell(row=2, column=5).value == "Empreendedor 1"
    assert ws.cell(row=2, column=7).value == "Empreendedorismo MEI"
    assert ws.cell(row=2, column=8).value == "30h"

    # Validar estilos: cabeçalho com fundo azul marinho (1E3A8A) e texto branco em negrito
    cell_a1 = ws.cell(row=1, column=1)
    assert cell_a1.font.bold is True
    assert cell_a1.font.color.rgb in ("00FFFFFF", "FFFFFF")
    assert cell_a1.fill.fill_type == "solid"
    assert "1E3A8A" in str(cell_a1.fill.start_color.rgb)


def test_custom_initial_sequence(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "custom_init.db"
    # Começando em livro 3, folha 99, registro 299
    manager = LivroRegistroManager(
        db_url_or_path=str(db_path),
        folhas_por_livro=100,
        initial_livro=3,
        initial_folha=99,
        initial_registro=299,
    )

    next_livro, next_folha, next_reg = manager.get_next_numbers()
    assert (next_livro, next_folha, next_reg) == (3, 99, 299)

    curso = CursoMetadata(curso_nome="Teste", carga_horaria=10, data_conclusao="01/01/2026")
    aluno1 = ValidacaoAluno("Aluno A", "111", "111", "***", True)
    aluno2 = ValidacaoAluno("Aluno B", "222", "222", "***", True)

    reg1 = manager.register_certificate(aluno1, curso)
    assert (reg1.livro_numero, reg1.folha_numero, reg1.registro_numero) == (3, 99, 299)

    reg2 = manager.register_certificate(aluno2, curso)
    assert (reg2.livro_numero, reg2.folha_numero, reg2.registro_numero) == (3, 100, 300)

    # Terceiro deve virar para Livro 4, Folha 1, Registro 301
    aluno3 = ValidacaoAluno("Aluno C", "333", "333", "***", True)
    reg3 = manager.register_certificate(aluno3, curso)
    assert (reg3.livro_numero, reg3.folha_numero, reg3.registro_numero) == (4, 1, 301)


def test_register_certificate_without_cpf(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "sem_cpf.db"
    excel_path = tmp_path / "livro_sem_cpf.xlsx"
    manager = LivroRegistroManager(
        db_url_or_path=str(db_path),
        excel_export_path=str(excel_path),
    )

    curso = CursoMetadata(
        curso_nome="Workshop Prático",
        carga_horaria=8,
        data_conclusao="10/10/2026",
    )
    aluno_sem_cpf = ValidacaoAluno(
        nome="Mariana Sem CPF",
        cpf="",
        cpf_formatado="",
        cpf_mascarado="",
        is_valido=True,
    )

    reg = manager.register_certificate(aluno_sem_cpf, curso)
    assert reg.aluno_nome == "Mariana Sem CPF"
    assert reg.aluno_cpf == ""
    assert reg.aluno_cpf_mascarado == "Não informado"
    assert len(reg.codigo_autenticidade) == 64

    # Busca por código deve funcionar perfeitamente
    buscado = manager.get_certificate_by_code(reg.codigo_autenticidade)
    assert buscado is not None
    assert buscado.aluno_nome == "Mariana Sem CPF"
    assert buscado.aluno_cpf_mascarado == "Não informado"

    # Excel exportado deve conter hífen no CPF
    assert excel_path.exists()
    wb = openpyxl.load_workbook(str(excel_path))
    ws = wb.active
    # Linha 2, coluna 6 (CPF)
    assert ws.cell(row=2, column=6).value == "-"


def test_update_certificate(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "update_test.db"
    excel_path = tmp_path / "update_test.xlsx"
    manager = LivroRegistroManager(db_url_or_path=str(db_path), excel_export_path=excel_path)

    curso = CursoMetadata(curso_nome="Teste Retificação", carga_horaria=20, data_conclusao="10/10/2026")
    aluno = ValidacaoAluno("Nome Errado", "52998224725", "529.982.247-25", "***.982.247-**", True)
    reg = manager.register_certificate(aluno, curso)

    # Atualiza nome e CPF
    atualizado = manager.update_certificate(
        codigo_autenticidade=reg.codigo_autenticidade,
        novo_nome="Nome Corrigido Silva",
        novo_cpf="111.444.777-35",
    )

    assert atualizado.aluno_nome == "Nome Corrigido Silva"
    assert atualizado.aluno_cpf == "11144477735"
    assert atualizado.aluno_cpf_mascarado == "***.444.777-**"

    # Confere no banco
    consultado = manager.get_certificate_by_code(reg.codigo_autenticidade)
    assert consultado.aluno_nome == "Nome Corrigido Silva"
    assert consultado.aluno_cpf == "11144477735"

    # Rejeita CPF inválido
    with pytest.raises(ValueError, match="CPF inválido"):
        manager.update_certificate(reg.codigo_autenticidade, novo_cpf="111.111.111-11")

    # Rejeita nome sem sobrenome
    with pytest.raises(ValueError, match="sobrenome"):
        manager.update_certificate(reg.codigo_autenticidade, novo_nome="ApenasNome")


def test_reset_database(tmp_path):
    from core.registry import CursoMetadata, LivroRegistroManager

    db_path = tmp_path / "reset_test.db"
    excel_path = tmp_path / "reset_test.xlsx"
    manager = LivroRegistroManager(db_url_or_path=str(db_path), excel_export_path=excel_path)

    curso = CursoMetadata(curso_nome="Curso Teste", carga_horaria=10, data_conclusao="01/01/2026")
    aluno = ValidacaoAluno("Aluno Teste", "52998224725", "529.982.247-25", "***.982.247-**", True)
    manager.register_certificate(aluno, curso)
    assert len(manager.list_all_certificates()) == 1

    # Reset do banco com redefinição de numeração inicial
    manager.reset_database(initial_livro=2, initial_folha=10, initial_registro=110)
    assert len(manager.list_all_certificates()) == 0
    assert manager.get_next_numbers() == (2, 10, 110)
