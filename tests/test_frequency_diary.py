"""Testes do Diário de Frequência, histórico de turmas, retificação de faltas e emissão não-bloqueante."""

import io
from datetime import datetime
from pathlib import Path
import openpyxl
import pytest

from core.batch_service import emitir_lote_certificados
from core.registry import CursoMetadata, LivroRegistroManager
from core.renderer import CertificateRenderConfig
from core.validator import validar_aluno


@pytest.fixture
def manager(tmp_path: Path) -> LivroRegistroManager:
    """Instância isolada do LivroRegistroManager em banco SQLite temporário."""
    db_file = tmp_path / "test_registros.db"
    mgr = LivroRegistroManager(db_url_or_path=str(db_file))
    return mgr


def test_criar_turma_lote_e_obter_diario(manager: LivroRegistroManager):
    """Testa a criação de turma_lote, registro de presenças e consulta do diário."""
    curso = CursoMetadata(
        curso_nome="Curso Prático de IA",
        carga_horaria=40,
        data_inicio="01/09/2026",
        data_conclusao="10/09/2026",
        data_emissao="11/09/2026",
        instrutor="Prof. Alan Turing",
        cidade="Goiânia",
    )
    encontros = [
        {"nome_coluna": "(4h) 2026/set/01", "horas": 4, "data_str": "01/09/2026"},
        {"nome_coluna": "(4h) 2026/set/02", "horas": 4, "data_str": "02/09/2026"},
    ]
    lote_id = manager.criar_turma_lote(
        curso=curso,
        total_alunos=2,
        encontros=encontros,
        identificador="Turma IA 2026",
    )
    assert lote_id is not None
    assert lote_id > 0

    turmas = manager.obter_turmas_lotes()
    assert len(turmas) == 1
    assert turmas[0]["identificador_lote"] == "Turma IA 2026"
    assert turmas[0]["total_alunos"] == 2

    # Registra certificados vinculados ao lote
    a1 = validar_aluno(
        "Ada Lovelace",
        "52998224725",
        frequencia=100,
        detalhes_presenca={"(4h) 2026/set/01": True, "(4h) 2026/set/02": True},
    )
    manager.register_batch([a1], curso, lote_id=lote_id)

    # Consulta diário da turma
    diario = manager.obter_diario_turma(lote_id)
    assert diario["turma"]["id"] == lote_id
    assert len(diario["encontros"]) == 2
    assert len(diario["alunos"]) == 1
    assert diario["alunos"][0]["aluno_nome"] == "Ada Lovelace"
    assert diario["alunos"][0]["presencas"].get("(4h) 2026/set/01") is True


def test_retificar_frequencia_aluno(manager: LivroRegistroManager):
    """Testa retificação de frequência por motivo legal com registro em auditoria."""
    a1 = validar_aluno("Grace Hopper", "52998224725", frequencia=60)
    curso = CursoMetadata(curso_nome="Compiladores", carga_horaria=20, data_conclusao="10/09/2026")
    registros = manager.register_batch([a1], curso)
    cod = registros[0].codigo_autenticidade

    # Retifica frequência de 60% para 85% com atestado
    rec_atualizado = manager.retificar_frequencia_aluno(
        codigo_autenticidade=cod,
        nova_frequencia=85,
        justificativa="Atestado médico de 2 dias apresentado em 11/09/2026",
    )
    assert rec_atualizado.frequencia == 85

    # Confere no banco de dados
    rec_db = manager.get_certificate_by_code(cod)
    assert rec_db is not None
    assert rec_db.frequencia == 85

    # Confere se foi gravado na auditoria
    historico = manager.obter_historico_consultas()
    retif_logs = [h for h in historico if h["status"] == "RETIFICACAO_FREQUENCIA"]
    assert len(retif_logs) == 1
    assert "85%" in retif_logs[0]["aluno_nome"]
    assert "Atestado médico" in retif_logs[0]["curso_nome"]


def test_exportar_diario_classe_excel(manager: LivroRegistroManager):
    """Testa a geração e formatação do Diário de Classe em planilha Excel."""
    curso = CursoMetadata(
        curso_nome="Robótica Básica",
        carga_horaria=20,
        data_inicio="01/09/2026",
        data_conclusao="05/09/2026",
        data_emissao="06/09/2026",
        instrutor="Prof. Nikola Tesla",
        cidade="Goiânia",
    )
    encontros = [{"nome_coluna": "(4h) 2026/set/01", "horas": 4, "data_str": "01/09/2026"}]
    lote_id = manager.criar_turma_lote(
        curso=curso,
        total_alunos=1,
        encontros=encontros,
        identificador="Turma Robótica 2026",
    )

    a1 = validar_aluno(
        "Marie Curie",
        "52998224725",
        frequencia=100,
        detalhes_presenca={"(4h) 2026/set/01": True},
    )
    curso = CursoMetadata(curso_nome="Robótica Básica", carga_horaria=20, data_conclusao="05/09/2026")
    manager.register_batch([a1], curso, lote_id=lote_id)

    excel_buf = io.BytesIO()
    manager.exportar_diario_classe_excel(
        lote_id=lote_id,
        instituicao_nome="Escola do Futuro",
        razao_social="Instituto Futuro Facil",
        cnpj="00.000.000/0001-91",
        destination=excel_buf,
    )
    excel_buf.seek(0)

    wb = openpyxl.load_workbook(excel_buf)
    ws = wb.active
    assert ws is not None
    assert ws.title == "Diário de Classe"
    # Cabeçalho da instituição
    assert "ESCOLA DO FUTURO" in str(ws.cell(row=1, column=1).value)
    # Linha do aluno
    assert ws.cell(row=6, column=2).value == "Marie Curie"
    # Frequência e Situação
    assert ws.cell(row=6, column=6).value == "100%"
    assert ws.cell(row=6, column=7).value == "APROVADO"


def test_cpf_resiliente_zfill():
    """Testa recuperação de CPF com zeros removidos pelo Excel."""
    # CPF válido 012.345.678-90 -> 10 dígitos "1234567890"
    aluno = validar_aluno("Santos Dumont", "1234567890")
    assert aluno.cpf == "01234567890"
    assert aluno.is_valido is True
    assert aluno.cpf_mascarado == "***.345.678-**"


def test_emissao_lote_com_frequencia_e_diario(manager: LivroRegistroManager):
    """Testa emissão de lote com metadados de encontros e verificação de pacote ZIP gerado."""
    a1 = validar_aluno(
        "Carlos Drummond de Andrade",
        "52998224725",
        frequencia=100,
        detalhes_presenca={"(4h) 2026/set/01": True},
    )
    curso = CursoMetadata(
        curso_nome="Literatura Moderna",
        carga_horaria=40,
        data_conclusao="10/09/2026",
    )
    config = CertificateRenderConfig()
    encontros = [{"nome_coluna": "(4h) 2026/set/01", "horas": 4, "data_str": "01/09/2026"}]

    res = emitir_lote_certificados(
        alunos=[a1],
        curso=curso,
        config=config,
        manager=manager,
        encontros=encontros,
        identificador_turma="Literatura 2026",
    )
    assert res.total_emitidos == 1
    assert res.lote_id is not None
    assert len(res.zip_bytes) > 0

    # Verifica se a turma foi gravada com lote_id
    turmas = manager.obter_turmas_lotes()
    assert any(t["id"] == res.lote_id for t in turmas)


def test_validacao_regras_grill_me():
    """Testa detalhadamente as regras de negócio alinhadas no grill-me:
    1. Aluno com frequência < 75% não bloqueia emissão de outros alunos aptos.
    2. Aluno com frequência < 75% pode ser aprovado alterando manualmente a frequência para >= 75%.
    3. As faltas originais (detalhes_presenca) são preservadas intactas no diário.
    4. Nome monônimo (único) é estritamente bloqueante (tem_erro_cadastral=True).
    5. CPF inválido em modo não-obrigatório emite como sem CPF sem travar o lote.
    """
    presencas_originais = {"(4h) Encontro 1": True, "(4h) Encontro 2": False}

    # 1. Aluno reprovado por frequência (< 75%)
    aluno_rep = validar_aluno(
        "Joaquim José da Silva Xavier",
        "52998224725",
        frequencia=50,
        frequencia_minima=75,
        detalhes_presenca=presencas_originais,
    )
    assert aluno_rep.is_valido is False
    assert aluno_rep.tem_erro_frequencia is True
    assert aluno_rep.tem_erro_cadastral is False
    assert aluno_rep.is_apto_emissao is False

    # 2. Aluno aprovado manualmente para 75%
    aluno_aprovado_manual = validar_aluno(
        "Joaquim José da Silva Xavier",
        "52998224725",
        frequencia=75,
        frequencia_minima=75,
        detalhes_presenca=presencas_originais,
    )
    assert aluno_aprovado_manual.is_valido is True
    assert aluno_aprovado_manual.tem_erro_frequencia is False
    assert aluno_aprovado_manual.tem_erro_cadastral is False
    assert aluno_aprovado_manual.is_apto_emissao is True
    # 3. Faltas originais preservadas intactas
    assert aluno_aprovado_manual.detalhes_presenca == presencas_originais

    # 4. Monônimo é estritamente bloqueante
    aluno_mononimo = validar_aluno("Aristóteles", "52998224725", frequencia=100)
    assert aluno_mononimo.is_valido is False
    assert aluno_mononimo.tem_erro_cadastral is True
    assert aluno_mononimo.is_apto_emissao is False

    # 5. CPF inválido com fallback para sem CPF (não bloqueante)
    aluno_cpf_invalido_ignorado = validar_aluno(
        "Platão de Atenas",
        "11111111111",
        cpf_obrigatorio=False,
        permitir_cpf_invalido_como_sem_cpf=True,
    )
    assert aluno_cpf_invalido_ignorado.is_valido is True
    assert aluno_cpf_invalido_ignorado.tem_erro_cadastral is False
    assert aluno_cpf_invalido_ignorado.cpf_invalido_ignorado is True
    assert aluno_cpf_invalido_ignorado.cpf == ""
    assert aluno_cpf_invalido_ignorado.cpf_formatado == ""

