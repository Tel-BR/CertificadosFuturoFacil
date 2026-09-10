import io
from pathlib import Path
import openpyxl
import pandas as pd
import pytest

from core.spreadsheet import (
    generate_template_spreadsheet,
    read_and_validate_spreadsheet,
    SpreadsheetValidationResult,
)


def test_generate_template_to_path(tmp_path: Path):
    target_file = tmp_path / "modelo_alunos.xlsx"
    res_path = generate_template_spreadsheet(target_file)
    assert Path(res_path).exists()

    wb = openpyxl.load_workbook(target_file)
    sheet = wb.active
    assert sheet is not None

    headers = [cell.value for cell in sheet[1]]
    assert headers[:2] == ["Nome", "CPF"]

    # Modelo deve conter estritamente o cabeçalho (sem dados de teste residuais)
    assert sheet.max_row == 1
    wb.close()


def test_generate_template_to_bytes():
    buffer = io.BytesIO()
    generate_template_spreadsheet(buffer)
    buffer.seek(0)

    wb = openpyxl.load_workbook(buffer)
    sheet = wb.active
    assert sheet is not None
    headers = [cell.value for cell in sheet[1]]
    assert headers[:2] == ["Nome", "CPF"]
    assert sheet.max_row == 1
    wb.close()


def test_read_and_validate_valid_spreadsheet(tmp_path: Path):
    test_file = tmp_path / "alunos_validos.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["MARIA DA SILVA", "joão dos santos"],
            "CPF": ["529.982.247-25", "111.444.777-35"],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert isinstance(result, SpreadsheetValidationResult)
    assert result.is_valid is True
    assert result.total_rows == 2
    assert result.valid_count == 2
    assert result.invalid_count == 0
    assert len(result.alunos) == 2
    assert result.alunos[0].nome == "Maria da Silva"
    assert result.alunos[0].cpf_formatado == "529.982.247-25"
    assert result.alunos[1].nome == "João dos Santos"


def test_read_and_validate_with_invalid_rows(tmp_path: Path):
    test_file = tmp_path / "alunos_misto.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Carlos Silva", "   ", "Ana Clara"],
            "CPF": ["529.982.247-25", "11144477735", "123.456.789-00"],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.is_valid is False
    assert result.total_rows == 3
    assert result.valid_count == 1
    assert result.invalid_count == 2
    assert result.alunos[0].is_valido is True
    assert result.alunos[1].is_valido is False  # Nome vazio
    assert result.alunos[2].is_valido is False  # CPF inválido


def test_read_and_validate_missing_columns(tmp_path: Path):
    test_file = tmp_path / "colunas_erradas.xlsx"
    df = pd.DataFrame(
        {
            "Participante": ["Carlos Silva"],
            "Documento": ["529.982.247-25"],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.is_valid is False
    assert len(result.global_errors) > 0
    assert any("Nome" in err or "CPF" in err for err in result.global_errors)


def test_read_and_validate_rejects_extra_columns(tmp_path: Path):
    # Requisito: estritamente duas colunas (Nome e CPF)
    test_file = tmp_path / "colunas_extras.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Carlos Silva"],
            "CPF": ["529.982.247-25"],
            "Email": ["carlos@exemplo.com"],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.is_valid is False
    assert any("não permitidas" in err for err in result.global_errors)


def test_read_and_validate_csv_with_semicolon(tmp_path: Path):
    test_file = tmp_path / "alunos_semicolon.csv"
    test_file.write_text("Nome;CPF\nLucas Souza;529.982.247-25\n", encoding="utf-8")

    result = read_and_validate_spreadsheet(test_file)
    assert result.is_valid is True
    assert result.total_rows == 1
    assert result.alunos[0].nome == "Lucas Souza"


def test_read_and_validate_numeric_cpf_handling(tmp_path: Path):
    test_file = tmp_path / "alunos_numericos.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Aluno Zero", "Aluno Curto"],
            # 1234567890 tem 10 dígitos (zero inicial foi comido pelo Excel -> 01234567890 válido)
            # 123 é curto demais e não deve sofrer padding artificial
            "CPF": [1234567890, 123],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.alunos[0].cpf == "01234567890"
    assert result.alunos[0].is_valido is True
    assert result.alunos[1].cpf == "123"
    assert result.alunos[1].is_valido is False


def test_read_and_validate_spreadsheet_with_empty_cpfs(tmp_path: Path):
    test_file = tmp_path / "alunos_com_sem_cpf.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Aluno Com CPF", "Aluno Sem CPF"],
            "CPF": ["529.982.247-25", ""],
        }
    )
    df.to_excel(test_file, index=False)

    # Por padrão (cpf_obrigatorio=False), ambos são válidos e sem_cpf_count é 1
    result = read_and_validate_spreadsheet(test_file)
    assert result.is_valid is True
    assert result.total_rows == 2
    assert result.valid_count == 2
    assert result.invalid_count == 0
    assert result.sem_cpf_count == 1
    assert result.alunos[0].cpf == "52998224725"
    assert result.alunos[1].cpf == ""
    assert result.alunos[1].is_valido is True

    # Se for exigido CPF obrigatório, a linha sem CPF é reprovada
    result_obrigatorio = read_and_validate_spreadsheet(test_file, cpf_obrigatorio=True)
    assert result_obrigatorio.is_valid is False
    assert result_obrigatorio.valid_count == 1
    assert result_obrigatorio.invalid_count == 1
    assert result_obrigatorio.alunos[1].is_valido is False


def test_read_and_validate_encounter_columns_and_frequency(tmp_path: Path):
    """Verifica apuração automática de carga horária e cálculo de frequência com colunas de encontros."""
    test_file = tmp_path / "alunos_frequencia.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Alice Silva", "Bruno Santos", "Carla Dias"],
            "CPF": ["529.982.247-25", "111.444.777-35", "222.333.444-55"],
            "(4h) I 2026/ago/08": [True, True, True],
            "(4h) II 2026/ago/08": [True, True, False],
            "(2h) 2026/ago/09": [True, False, False],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.carga_horaria_calculada == 10  # 4 + 4 + 2 = 10h
    assert result.data_inicio_calculada == "08/08/2026"
    assert result.data_fim_calculada == "09/08/2026"
    assert len(result.encontros_detectados) == 3

    # Alice: 10/10h = 100% -> aprovada
    assert result.alunos[0].horas_presentes == 10
    assert result.alunos[0].frequencia == 100
    assert result.alunos[0].is_valido is True

    # Bruno: 8/10h = 80% (>= 75%) -> aprovado
    assert result.alunos[1].horas_presentes == 8
    assert result.alunos[1].frequencia == 80
    assert result.alunos[1].is_valido is True

    # Carla: 4/10h = 40% (< 75%) -> reprovada por falta
    assert result.alunos[2].horas_presentes == 4
    assert result.alunos[2].frequencia == 40
    assert result.alunos[2].is_valido is False
    assert any("75%" in err for err in result.alunos[2].erros)

    assert result.is_valid is False  # pois Carla foi reprovada
    assert result.valid_count == 2
    assert result.invalid_count == 1


def test_generate_template_with_dates_and_example(tmp_path: Path):
    """Testa geração de planilha modelo dinâmica baseada nas datas da turma com coluna de aproveitamento."""
    target_file = tmp_path / "modelo_turma.xlsx"
    generate_template_spreadsheet(
        target_file,
        data_inicio="10/09/2026",
        data_fim="15/09/2026",
        horas_por_encontro=4,
        incluir_exemplo=True,
    )
    assert target_file.exists()

    wb = openpyxl.load_workbook(target_file)
    ws = wb.active
    assert ws is not None

    headers = [cell.value for cell in ws[1]]
    assert headers[0] == "Nome"
    assert headers[1] == "CPF"
    assert headers[2] == "Aproveitamento (%)"
    assert headers[3] == "(4h) 2026/set/10"
    assert headers[4] == "(4h) 2026/set/15"

    assert ws.max_row == 2
    exemplo_row = [cell.value for cell in ws[2]]
    assert exemplo_row[0] == "Maria de Souza Silva"
    assert exemplo_row[1] == "529.982.247-25"
    assert exemplo_row[2] is None
    assert exemplo_row[3] is True
    assert exemplo_row[4] is False
    wb.close()


def test_manual_aproveitamento_precedence_over_encounters(tmp_path: Path):
    """
    Testa se o número digitado na coluna de aproveitamento se torna a fonte da verdade,
    dispensando o cálculo por booleanos, e se quando vazia calcula normalmente pelos encontros.
    """
    test_file = tmp_path / "teste_precedencia.xlsx"
    df = pd.DataFrame(
        {
            "Nome": ["Aluno Um Silva", "Aluno Dois Santos", "Aluno Tres Costa"],
            "CPF": ["529.982.247-25", "111.444.777-35", "012.345.678-90"],
            "Aproveitamento (%)": [85, None, "100%"],  # Aluno 1: manual 85%, Aluno 2: vazio (calcula), Aluno 3: manual 100%
            "(4h) 2026/set/01": [True, True, False],
            "(4h) 2026/set/10": [False, False, False],
        }
    )
    df.to_excel(test_file, index=False)

    result = read_and_validate_spreadsheet(test_file)
    assert result.total_rows == 3

    # Aluno 1: presença nos booleanos daria 4h/8h = 50%, mas como forneceu 85%, a fonte da verdade é 85%!
    assert result.alunos[0].frequencia == 85
    assert result.alunos[0].is_valido is True

    # Aluno 2: aproveitamento manual estava vazio (None), então calcula normalmente pelos encontros (4h/8h = 50%)
    assert result.alunos[1].frequencia == 50
    assert result.alunos[1].is_valido is False  # reprovado (< 75%)

    # Aluno 3: presença nos booleanos daria 0h/8h = 0%, mas como forneceu 100%, a fonte da verdade é 100%!
    assert result.alunos[2].frequencia == 100
    assert result.alunos[2].is_valido is True



